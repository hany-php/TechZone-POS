import openpyxl
import sqlite3
import shutil
import datetime
import os
import pathlib

BASE_DIR = pathlib.Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "data" / "shop.db"
EXCEL_PATH = BASE_DIR / "المنتج (product.template) (1).xlsx"

# 1. Create Timestamped Backup
now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = BASE_DIR / "data" / f"shop_backup_{now_str}.db"
shutil.copyfile(DB_PATH, backup_path)
print(f"✓ Backup created successfully: {backup_path.name}")

# 2. Connect to database
conn = sqlite3.connect(DB_PATH)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

print("\n--- Current Data Before Reset ---")
for t in ["users", "customers", "settings", "categories", "products", "sales", "shifts", "repairs", "used_device_purchases", "quotations"]:
    cur.execute(f"SELECT COUNT(*) FROM {t}")
    print(f"  {t}: {cur.fetchone()[0]}")

# 3. Clean transaction tables and old products/categories
cleared_tables = [
    "serials", "repair_parts", "repairs", "used_device_purchases",
    "quotation_items", "quotations", "sale_return_items", "sale_returns",
    "payments", "customer_payments", "sale_items", "sales", "shifts",
    "stock_movements", "products", "categories"
]

for t in cleared_tables:
    cur.execute(f"DELETE FROM {t}")

cur.execute(
    f"DELETE FROM sqlite_sequence WHERE name IN ({','.join('?' for _ in cleared_tables)})",
    cleared_tables
)
print("✓ Previous transactions, mock products, and categories cleared.")

# 4. Clean old receipt PDFs in data/receipts/ if any
receipts_dir = BASE_DIR / "data" / "receipts"
if receipts_dir.exists():
    for f in receipts_dir.glob("*.pdf"):
        try:
            f.unlink()
        except Exception:
            pass
    print("✓ Old receipt PDFs cleared.")

# 5. Read and Import from Excel
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
sheet = wb.active

categories_map = {}
inserted_count = 0

for r in range(2, sheet.max_row + 1):
    name = sheet.cell(r, 1).value
    sku = sheet.cell(r, 2).value
    price_val = sheet.cell(r, 3).value
    cost_val = sheet.cell(r, 4).value
    cat_name = sheet.cell(r, 5).value
    qty_val = sheet.cell(r, 6).value
    img_val = sheet.cell(r, 7).value

    clean_name = str(name).strip() if name else f"منتج {r}"
    clean_cat = str(cat_name).strip() if cat_name and str(cat_name).strip() else "أخرى"

    if clean_cat not in categories_map:
        cur.execute("INSERT INTO categories(name) VALUES (?)", (clean_cat,))
        categories_map[clean_cat] = cur.lastrowid

    cat_id = categories_map[clean_cat]

    if sku and str(sku).strip():
        clean_sku = str(sku).strip()
    else:
        clean_sku = f"TZ-AUTO-{r:03d}"

    try:
        p_flt = float(price_val or 0)
    except Exception:
        p_flt = 0.0
    price_cents = int(round(p_flt * 100))

    try:
        c_flt = float(cost_val or 0)
    except Exception:
        c_flt = 0.0
    cost_cents = int(round(c_flt * 100))

    try:
        qty_int = int(qty_val or 0)
    except Exception:
        qty_int = 0
    qty_final = max(0, qty_int)

    if img_val and str(img_val).strip():
        data_url = f"data:image/webp;base64,{str(img_val).strip()}"
    else:
        data_url = None

    cur.execute(
        """
        INSERT INTO products(sku, name, category_id, price, cost, qty, min_qty, tracked, image, notes, active, created_at)
        VALUES (?, ?, ?, ?, ?, ?, 5, 0, ?, NULL, 1, datetime('now', 'localtime'))
        """,
        (clean_sku, clean_name, cat_id, price_cents, cost_cents, qty_final, data_url)
    )
    inserted_count += 1

conn.commit()

print(f"\n✓ Successfully imported {inserted_count} real products into {len(categories_map)} categories.")

print("\n--- Data After Reset & Import ---")
for t in ["users", "customers", "settings", "categories", "products", "sales", "shifts", "repairs", "used_device_purchases", "quotations"]:
    cur.execute(f"SELECT COUNT(*) FROM {t}")
    print(f"  {t}: {cur.fetchone()[0]}")

print("\n--- Categories Breakdown ---")
for c_name, c_id in sorted(categories_map.items(), key=lambda x: x[1]):
    cur.execute("SELECT COUNT(*) FROM products WHERE category_id=?", (c_id,))
    p_cnt = cur.fetchone()[0]
    print(f"  [{c_id}] {c_name}: {p_cnt} منتج")

conn.close()
print("\n✓ Database migration completed 100% successfully!")
